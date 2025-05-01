
import streamlit as st
import pandas as pd
from urllib.parse import urlparse
from io import BytesIO
from openpyxl.styles import Alignment
from openpyxl import load_workbook
import streamlit.components.v1 as components

st.set_page_config(page_title="AEM URL Converter", page_icon="🌍", layout="wide")

st.markdown("<h1 style='text-align: center; color: #2E86C1;'>🌍 AEM URL Converter</h1>", unsafe_allow_html=True)
st.markdown("Upload a Web Translation Excel file to convert AEM URLs based on target languages.")

LANGUAGE_MAP = {
    "de-DE": "/content/lifetech/europe/en-de",
    "es-ES": "/content/lifetech/europe/en-es",
    "fr-FR": "/content/lifetech/europe/en-fr",
    "ja-JP": "/content/lifetech/japan/en-jp",
    "ko-KR": "/content/lifetech/ipac/en-kr",
    "zh-CN": "/content/lifetech/greater-china/en-cn",
    "zh-TW": "/content/lifetech/ipac/en-tw",  # ✅ corrected
    "pt-BR": "/content/lifetech/latin-america/en-br",
    "es-LATAM": "/content/lifetech/latin-america/en-mx"
}

def clean_url(url):
    if not isinstance(url, str):
        return None
    parsed = urlparse(url)
    path = parsed.path
    if "/home/" not in path:
        return None
    cleaned = path.split("/home/", 1)[1]
    if ".html" in cleaned:
        cleaned = cleaned.replace(".html", "")
    return "/home/" + cleaned

def detect_first_url(row):
    for cell in row:
        if isinstance(cell, str) and "/home/" in cell:
            return cell
    return None

def process_file(uploaded_file):
    df = pd.read_excel(uploaded_file, sheet_name=0, header=3)
    results = []

    language_columns = {col: code for code, path in LANGUAGE_MAP.items() for col in df.columns if code in str(col)}

    for _, row in df.iterrows():
        original_url = detect_first_url(row)
        cleaned_path = clean_url(original_url)
        if not cleaned_path:
            continue

        for col_name, lang_code in language_columns.items():
            cell_value = row.get(col_name, "")
            if pd.notna(cell_value) and str(cell_value).strip().lower() not in ["", "no"]:
                localized_base = LANGUAGE_MAP.get(lang_code)
                if localized_base:
                    results.append({
                        "Original URL": original_url,
                        "Language": lang_code,
                        "Localized Path": localized_base + cleaned_path
                    })

    result_df = pd.DataFrame(results)
    return result_df.sort_values(by=["Language"]) if not result_df.empty else pd.DataFrame()

uploaded_file = st.file_uploader("📂 Upload Excel File", type=["xlsx"])

if uploaded_file:
    try:
        df_result = process_file(uploaded_file)
        if df_result.empty:
            st.warning("⚠️ No valid data found in the uploaded file.")
        else:
            st.success("✅ URLs converted successfully!")

            st.markdown("### 🔗 Localized URLs")
            for idx, row in df_result.iterrows():
                cols = st.columns([1, 5, 1])
                cols[0].markdown(f"**{row['Language']}**")
                cols[1].code(row['Localized Path'], language="bash")
                copy_script = f"""
                <button onclick="navigator.clipboard.writeText('{row['Localized Path']}')"
                        style="padding:5px 10px; background-color:#2E86C1; color:white; border:none; border-radius:5px;">
                    📋 Copy
                </button>
                """
                cols[2].components.html(copy_script, height=40)

            output = BytesIO()
            with pd.ExcelWriter(output, engine="openpyxl") as writer:
                df_result.to_excel(writer, index=False)
                worksheet = writer.sheets["Sheet1"]
                worksheet.column_dimensions["A"].width = 60
                worksheet.column_dimensions["B"].width = 20
                worksheet.column_dimensions["C"].width = 80
                for row in worksheet.iter_rows():
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                for cell in worksheet[1]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                for row in worksheet.iter_rows(min_row=2, min_col=2, max_col=2):
                    for cell in row:
                        cell.alignment = Alignment(horizontal="center", vertical="top")

            st.download_button(
                label="📥 Download Converted Excel",
                data=output.getvalue(),
                file_name="converted_urls.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
    except Exception as e:
        st.error(f"❌ An error occurred: {e}")
